"""Fill the corporate «Цели на ИС» Excel workbook from a role goals template."""
from __future__ import annotations

import base64
import datetime as _dt
import io
import pathlib
from openpyxl import load_workbook

from plan import load_template, goals_from_template

TEMPLATES_DIR = pathlib.Path(__file__).parent / "templates"
GOALS_TEMPLATE = TEMPLATES_DIR / "probation_goals.xlsx"
MAX_GOAL_ROWS = 10  # rows 11..20 on sheet «Цели»
GOALS_SHEET = "Цели"


def _add_months(d: _dt.date, months: int) -> _dt.date:
    y = d.year + (d.month - 1 + months) // 12
    m = (d.month - 1 + months) % 12 + 1
    # clamp day for shorter months
    for day in (d.day, 28, 27, 26, 25, 24, 23, 22, 21, 20, 19, 18, 17, 16, 15, 14, 13, 12, 11, 10, 9, 8, 7, 6, 5, 4, 3, 2, 1):
        try:
            return _dt.date(y, m, day)
        except ValueError:
            continue
    return _dt.date(y, m, 1)


def build_goals_preview(
    *,
    role: str,
    hire_id: str,
    start_date: str,
    team: str = "",
    manager_id: str = "",
) -> dict:
    """Structured preview for chat (table/text) — no file bytes."""
    tpl = load_template(role)
    start = _dt.date.fromisoformat(start_date)
    months = int(tpl.get("probation_months", 3))
    end = _add_months(start, months)
    goals = goals_from_template(tpl, start_date)
    weight_sum = round(sum(g["weight"] for g in goals), 4)
    return {
        "template": "probation_goals.xlsx",
        "source_form": "Задачи на ИС (обновление).xlsx",
        "hire_id": hire_id,
        "role": role,
        "team": team,
        "manager_id": manager_id,
        "start_date": start.isoformat(),
        "end_date": end.isoformat(),
        "probation_months": months,
        "checkpoints": [
            {"label": "6-ая неделя", "date": (start + _dt.timedelta(days=42)).isoformat(),
             "title": "Промежуточная встреча по результатам задач на ИС"},
            {"label": "12-ая неделя", "date": (start + _dt.timedelta(days=84)).isoformat(),
             "title": "Итоговая встреча по результатам задач на ИС"},
        ],
        "goals": goals,
        "weight_sum": weight_sum,
        "weight_ok": abs(weight_sum - 1.0) < 0.02,
        "pass_threshold_note": "Испытательный срок считается пройденным при выполнении целей ≥ 90% (см. лист «Комментарии»).",
    }


def fill_goals_workbook(
    *,
    role: str,
    hire_id: str,
    start_date: str,
    team: str = "",
    manager_id: str = "",
) -> bytes:
    """Return .xlsx bytes filled from the corporate goals template."""
    if not GOALS_TEMPLATE.exists():
        raise FileNotFoundError(f"Missing template: {GOALS_TEMPLATE}")

    preview = build_goals_preview(
        role=role, hire_id=hire_id, start_date=start_date, team=team, manager_id=manager_id,
    )
    if len(preview["goals"]) > MAX_GOAL_ROWS:
        raise ValueError(f"Template has {len(preview['goals'])} goals; max {MAX_GOAL_ROWS} for Excel form")

    wb = load_workbook(GOALS_TEMPLATE)
    ws = wb[GOALS_SHEET]

    ws["B1"] = hire_id  # opaque id — no PII FIO in MVP
    ws["B2"] = role
    ws["B3"] = team
    ws["B4"] = preview["start_date"]
    ws["B5"] = preview["end_date"]
    ws["B6"] = manager_id or ""

    # Clear / fill goal rows 11..20
    for i in range(MAX_GOAL_ROWS):
        row = 11 + i
        if i < len(preview["goals"]):
            g = preview["goals"][i]
            ws[f"A{row}"] = g["summary"]
            ws[f"B{row}"] = g["expected_result"]
            ws[f"C{row}"] = g["due"]
            ws[f"D{row}"] = g["weight"]
            ws[f"E{row}"] = None  # % выполнения — заполняет руководитель позже
            # keep F formula if present
            if ws[f"F{row}"].value is None:
                ws[f"F{row}"] = f'=IF(D{row}="","",E{row}*D{row})'
        else:
            ws[f"A{row}"] = None
            ws[f"B{row}"] = None
            ws[f"C{row}"] = None
            ws[f"D{row}"] = None
            ws[f"E{row}"] = None

    # Checkpoint dates (B23:F23 / B24:F24 are merged — write top-left B)
    ws["B23"] = f"{preview['checkpoints'][0]['title']} ({preview['checkpoints'][0]['date']})"
    ws["B24"] = f"{preview['checkpoints'][1]['title']} ({preview['checkpoints'][1]['date']})"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_goals_workbook_payload(
    *,
    role: str,
    hire_id: str,
    start_date: str,
    team: str = "",
    manager_id: str = "",
    include_base64: bool = True,
) -> dict:
    preview = build_goals_preview(
        role=role, hire_id=hire_id, start_date=start_date, team=team, manager_id=manager_id,
    )
    raw = fill_goals_workbook(
        role=role, hire_id=hire_id, start_date=start_date, team=team, manager_id=manager_id,
    )
    filename = f"Цели_ИС_{hire_id}_{preview['start_date']}.xlsx"
    out = {
        "ok": True,
        "preview": preview,
        "filename": filename,
        "size_bytes": len(raw),
        "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    if include_base64:
        out["content_base64"] = base64.b64encode(raw).decode("ascii")
    return out
